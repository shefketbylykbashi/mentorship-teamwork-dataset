"""Source-data sanity checks for Type 6 profile assumptions."""

from __future__ import annotations

import argparse
import csv
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, Mapping

from benchmark_generation.core.config import load_config
from benchmark_generation.type6.download_mslib import utc_now
from benchmark_generation.type6.select_sources import DATASETS, DEFAULT_CONFIG, find_release_root


def skill_count_sanity_check(
    source_dir: str | Path,
    metadata_dir: str | Path,
    config: Mapping[str, Any],
) -> Dict[str, Any]:
    metadata_dir = Path(metadata_dir)
    release_root = find_release_root(source_dir)
    workbook_counts = workbook_skill_counts(release_root / "MSLIB (Parameters and BKS).xlsx")
    parsed_counts = parsed_candidate_skill_counts(metadata_dir / "type6_source_candidates.csv")
    raw_samples = raw_header_samples(release_root)
    datasets_pass = {}
    for dataset in DATASETS:
        workbook_unique = sorted(int(value) for value in workbook_counts.get(dataset, {}))
        parsed_unique = sorted(int(value) for value in parsed_counts.get(dataset, {}))
        expected = int(config["expected_counts"][dataset])
        datasets_pass[dataset] = (
            workbook_unique == [4]
            and parsed_unique == [4]
            and sum(workbook_counts[dataset].values()) == expected
            and sum(parsed_counts[dataset].values()) == expected
        )
    report = {
        "built_at": utc_now(),
        "official_documentation": {
            "MSRCPSP_DataFormat.pdf_present": (release_root / "MSRCPSP_DataFormat.pdf").exists(),
            "workbook_parameter_overview_definition": "# of skill types = The total number of skill types",
            "source_header_interpretation": (
                "First non-section row in each .msrcp Project Module is "
                "activities, resources, skill types, skill levels; the third field is n_skills."
            ),
        },
        "workbook_unique_skill_counts_per_set": {
            dataset: sorted(int(value) for value in workbook_counts.get(dataset, {}))
            for dataset in DATASETS
        },
        "workbook_skill_count_frequencies_per_set": {
            dataset: dict(sorted(workbook_counts.get(dataset, {}).items()))
            for dataset in DATASETS
        },
        "parsed_unique_skill_counts_per_set": {
            dataset: sorted(int(value) for value in parsed_counts.get(dataset, {}))
            for dataset in DATASETS
        },
        "parsed_skill_count_frequencies_per_set": {
            dataset: dict(sorted(parsed_counts.get(dataset, {}).items()))
            for dataset in DATASETS
        },
        "workbook_and_parser_agree_exactly": workbook_counts == parsed_counts,
        "raw_header_samples": raw_samples,
        "all_mslib1_4_sources_have_four_skills": all(datasets_pass.values()),
        "selected_100_four_skills_is_genuine_source_property": all(datasets_pass.values()),
        "dataset_checks": datasets_pass,
        "methodological_note": (
            "All canonical MSLIB1-4 source instances expose exactly four skill types. "
            "Type 6 must therefore preserve n_skills=4 and must not artificially "
            "diversify skill count for this family."
        ),
    }
    write_skill_count_sanity_outputs(report, metadata_dir)
    return report


def workbook_skill_counts(path: str | Path) -> Dict[str, Counter]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to inspect MSLIB workbook skill counts") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["MSRCPSP (basic)"]
    headers = [
        str(value) if value is not None else ""
        for value in next(ws.iter_rows(min_row=10, max_row=10, values_only=True))
    ]
    skill_idx = headers.index("# Skills")
    counts: Dict[str, Counter] = {dataset: Counter() for dataset in DATASETS}
    canonical = re.compile(r"^MSLIB_Set([1-4])_\d+\.msrcp$", re.IGNORECASE)
    for row in ws.iter_rows(min_row=11, values_only=True):
        dataset = str(row[0]) if row and row[0] is not None else ""
        filename = str(row[1]) if len(row) > 1 and row[1] is not None else ""
        if dataset in DATASETS and canonical.match(filename):
            counts[dataset][int(row[skill_idx])] += 1
    return counts


def parsed_candidate_skill_counts(path: str | Path) -> Dict[str, Counter]:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Missing parsed Type 6 source candidate table: {path}")
    counts: Dict[str, Counter] = {dataset: Counter() for dataset in DATASETS}
    with path.open(encoding="utf-8", newline="") as fh:
        for row in csv.DictReader(fh):
            dataset = str(row["MSLIB_set"])
            if dataset in DATASETS:
                counts[dataset][int(float(row["n_skills"]))] += 1
    return counts


def raw_header_samples(release_root: str | Path) -> Dict[str, Any]:
    release_root = Path(release_root)
    ids = {
        "MSLIB1": [1, 2322, 6600],
        "MSLIB2": [1, 1891, 9000],
        "MSLIB3": [1, 5863, 11880],
        "MSLIB4": [1, 3206, 5000],
    }
    samples = {}
    for dataset, source_ids in ids.items():
        rows = []
        with zipfile.ZipFile(release_root / f"{dataset}.zip") as zf:
            for source_id in source_ids:
                entry = f"{dataset}/MSLIB_Set{dataset[-1]}_{source_id}.msrcp"
                lines = zf.read(entry).decode("utf-8").splitlines()
                header = next(
                    line.strip()
                    for line in lines
                    if line.strip() and not line.startswith("\\*")
                )
                parts = [int(value) for value in header.split()]
                rows.append(
                    {
                        "source_zip_entry": entry,
                        "raw_header": header,
                        "activities": parts[0],
                        "resources": parts[1],
                        "skills": parts[2],
                        "skill_levels": parts[3],
                    }
                )
        samples[dataset] = rows
    return samples


def write_skill_count_sanity_outputs(report: Mapping[str, Any], metadata_dir: str | Path) -> None:
    metadata_dir = Path(metadata_dir)
    metadata_dir.mkdir(parents=True, exist_ok=True)
    (metadata_dir / "type6_skill_count_sanity_check.json").write_text(
        json.dumps(report, indent=2, sort_keys=True, default=str) + "\n",
        encoding="utf-8",
    )
    rows = []
    for dataset in DATASETS:
        rows.append(
            {
                "MSLIB_set": dataset,
                "workbook_unique_skill_counts": json.dumps(report["workbook_unique_skill_counts_per_set"][dataset]),
                "parsed_unique_skill_counts": json.dumps(report["parsed_unique_skill_counts_per_set"][dataset]),
                "workbook_frequencies": json.dumps(report["workbook_skill_count_frequencies_per_set"][dataset], sort_keys=True),
                "parsed_frequencies": json.dumps(report["parsed_skill_count_frequencies_per_set"][dataset], sort_keys=True),
                "dataset_check_passed": report["dataset_checks"][dataset],
            }
        )
    with (metadata_dir / "type6_skill_count_sanity_check.csv").open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=sorted(rows[0]))
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", default="type6_source_data")
    parser.add_argument("--metadata-dir", default="benchmark_metadata")
    parser.add_argument("--config", default="benchmark_generation/configs/type6.yaml")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    config = load_config(args.config, DEFAULT_CONFIG)
    report = skill_count_sanity_check(args.source_dir, args.metadata_dir, config)
    print(json.dumps({
        "all_mslib1_4_sources_have_four_skills": report["all_mslib1_4_sources_have_four_skills"],
        "workbook_and_parser_agree_exactly": report["workbook_and_parser_agree_exactly"],
        "selected_100_four_skills_is_genuine_source_property": report["selected_100_four_skills_is_genuine_source_property"],
    }, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
